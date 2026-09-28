"""Extract plain text from uploaded files and web pages."""

import csv
import io
import re
from html.parser import HTMLParser

import requests
from django.conf import settings

from common.netsafety import assert_public_url

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}


class LoaderError(Exception):
    pass


def load_file(django_file, filename):
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    django_file.open("rb")
    try:
        data = django_file.read()
    finally:
        django_file.close()
    if ext == ".pdf":
        return _load_pdf(data)
    if ext == ".docx":
        return _load_docx(data)
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    if ext == ".csv":
        return _load_csv(data.decode("utf-8", errors="replace"))
    raise LoaderError(f"Unsupported file type {ext or '(none)'}.")


def _load_pdf(data):
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:  # noqa: BLE001 - pypdf raises many types
        raise LoaderError(f"Could not read PDF: {exc}") from exc


def _load_docx(data):
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise LoaderError(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _load_csv(text):
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    header, lines = rows[0], []
    for row in rows[1:]:
        lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False) if v))
    return "\n".join(lines)


class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header"}
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article"}

    def __init__(self):
        super().__init__()
        self.parts, self._skip_depth = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def load_url(url):
    assert_public_url(url)
    try:
        response = requests.get(
            url, timeout=15, allow_redirects=False,
            headers={"User-Agent": "WhatsAppAISaaS-KnowledgeBot/1.0"},
        )
    except requests.RequestException as exc:
        raise LoaderError(f"Could not fetch URL: {exc.__class__.__name__}") from exc
    if response.status_code != 200:
        raise LoaderError(f"URL returned HTTP {response.status_code}.")
    if len(response.content) > settings.KNOWLEDGE_MAX_UPLOAD_BYTES:
        raise LoaderError("Page is too large.")
    content_type = response.headers.get("Content-Type", "")
    if "html" not in content_type:
        return response.text
    parser = _TextExtractor()
    parser.feed(response.text)
    return re.sub(r"\n\s*\n+", "\n\n", "".join(parser.parts)).strip()
